"""Salary sheet upload: mark a month's salary paid from the sheet HR already makes.

At month end HR uploads the signed salary sheet: the Excel file (read exactly)
or a PDF / photo (read by Gemini). Each row is matched to an employee by Staff
ID and shown next to what the system has for that month. Nothing changes until
HR confirms. On confirm, every ticked row becomes Paid with the sheet's Total
Salary as the final amount; the system's own figure and the sheet row are kept
in the payroll history, and the uploaded file is kept with the import.

An import can be undone for 48 hours by an Admin: every record goes back to
exactly what it was before.
"""
import base64
import io
import json
import logging
import re
import time
from html import escape

from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from sqlalchemy import text

from app.database import get_db

logger = logging.getLogger(__name__)
router = APIRouter()

MAX_FILE_BYTES = 12 * 1024 * 1024
UNDO_HOURS = 48
FIELDS = ("basic", "night", "friday", "present", "absent", "leave", "gross", "absent_deduction",
          "overtime", "extra", "net", "total")
MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august",
          "september", "october", "november", "december"]


def apply_salary_sheet_migrations(engine, sqlite: bool) -> None:
    pk = "INTEGER PRIMARY KEY AUTOINCREMENT" if sqlite else "BIGSERIAL PRIMARY KEY"
    big = "INTEGER" if sqlite else "BIGINT"
    with engine.begin() as conn:
        conn.execute(text(f"""CREATE TABLE IF NOT EXISTS salary_sheet_imports(
            id {pk}, salary_month TEXT NOT NULL, file_name TEXT NOT NULL, file_type TEXT NOT NULL,
            file_data TEXT NOT NULL, source TEXT NOT NULL, rows_json TEXT NOT NULL,
            printed_total DOUBLE PRECISION, month_label TEXT NOT NULL DEFAULT '',
            created_by TEXT NOT NULL DEFAULT '', created_at {big} NOT NULL,
            applied_at {big}, applied_by TEXT, results_json TEXT, undone_at {big}, undone_by TEXT)"""))


# ------------------------------------------------------------------ reading

def _number(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    digits = re.sub(r"[^\d.\-]", "", str(value).translate(str.maketrans("০১২৩৪৫৬৭৮৯", "0123456789")))
    try:
        return float(digits) if digits not in ("", "-", ".") else None
    except ValueError:
        return None


def _column_for(header: str) -> str | None:
    h = re.sub(r"\s+", " ", str(header or "")).strip().lower()
    if not h:
        return None
    if "staff" in h or h in ("id", "staff id"):
        return "staff_id"
    if h == "name" or h.endswith(" name") and "staff" not in h:
        return "name"
    if "design" in h or "esignat" in h:
        return "designation"
    if "absent" in h and ("deduc" in h or "duduc" in h or "duction" in h):
        return "absent_deduction"
    for key, words in (("basic", ("basic",)), ("night", ("night",)), ("friday", ("friday",)), ("present", ("present",)),
                       ("absent", ("absent",)), ("leave", ("leave",)), ("gross", ("gross", "grs", "rss sal")),
                       ("overtime", ("over", "ot")), ("extra", ("extra", "bonus")), ("net", ("net",)), ("total", ("total",))):
        if any(h == w or h.startswith(w) or f" {w}" in f" {h}" for w in words):
            return key
    return None


def read_excel(content: bytes) -> dict:
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("Excel file খোলা যায়নি। .xlsx file দিন।") from exc
    sheet = book.worksheets[0]
    table = [list(row) for row in sheet.iter_rows(values_only=True)]
    month_label = ""
    header_at, columns = None, {}
    for index, row in enumerate(table[:40]):
        joined = " ".join(str(v) for v in row if v is not None)
        found = re.search(r"month\s*[:\-]\s*([A-Za-z]+)", joined, re.I)
        if found and not month_label:
            month_label = found.group(1)
        mapped = {i: _column_for(v) for i, v in enumerate(row)}
        if "staff_id" in mapped.values() and ("total" in mapped.values() or "net" in mapped.values()):
            header_at, columns = index, {}
            for i, key in mapped.items():
                if key and key not in columns:
                    columns[key] = i
            break
    if header_at is None:
        raise ValueError("Sheet-এ 'Staff ID' আর 'Total Salary' (বা 'Net Salary') কলাম পাওয়া যায়নি।")
    rows, printed_total = [], None
    for row in table[header_at + 1:]:
        cell = lambda key: row[columns[key]] if key in columns and columns[key] < len(row) else None
        staff = str(cell("staff_id") or "").strip()
        values = {key: _number(cell(key)) for key in FIELDS}
        if not staff:
            total = values.get("total") if values.get("total") is not None else values.get("net")
            if total and not str(cell("name") or "").strip():
                printed_total = total                          # the yellow grand-total row
            continue
        rows.append({"staff_id": staff, "name": str(cell("name") or "").strip(),
                     "designation": str(cell("designation") or "").strip(), **values})
    return {"rows": rows, "printed_total": printed_total, "month_label": month_label}


AI_SYSTEM = (
    "You read a company's monthly salary sheet (a PDF or a photo, maybe rotated). Copy numbers exactly as printed; "
    "never calculate, round or guess. If a cell cannot be read, use null. Reply with JSON only.")
AI_PROMPT = (
    'Return {"month_label": "<month written on the sheet>", "printed_total": <grand total of the Total Salary column or null>, '
    '"rows": [{"name": "", "designation": "", "staff_id": "", "basic": 0, "night": 0, "friday": 0, "present": 0, '
    '"absent": 0, "leave": 0, "gross": 0, "absent_deduction": 0, "overtime": 0, "extra": 0, "net": 0, "total": 0}]}. '
    "One object per employee row, top to bottom. gross = the gross salary column, net = Net Salary, total = Total Salary. "
    "Skip the header and the grand-total row.")


def read_with_ai(content: bytes, mime: str) -> dict:
    from app.ai_insights import ask_gemini
    answer = ask_gemini(AI_SYSTEM, AI_PROMPT, files=[(mime, content)], temperature=0)
    if not isinstance(answer, dict) or not isinstance(answer.get("rows"), list):
        raise ValueError("AI sheet-টি পড়তে পারেনি। আরও পরিষ্কার ছবি বা PDF দিন, অথবা Excel দিন।")
    rows = []
    for item in answer["rows"]:
        if not isinstance(item, dict) or not str(item.get("staff_id") or "").strip():
            continue
        rows.append({"staff_id": str(item["staff_id"]).strip(), "name": str(item.get("name") or "").strip(),
                     "designation": str(item.get("designation") or "").strip(),
                     **{key: _number(item.get(key)) for key in FIELDS}})
    return {"rows": rows, "printed_total": _number(answer.get("printed_total")),
            "month_label": str(answer.get("month_label") or "")[:30]}


# ---------------------------------------------------------------- matching

def _norm(staff_id: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(staff_id or "").upper())


def _digits(staff_id: str) -> str:
    return re.sub(r"\D", "", str(staff_id or "")).lstrip("0")


def plan(rows: list[dict], month: str) -> list[dict]:
    """Match each sheet row to an employee and say what confirming will do."""
    with get_db() as c:
        employees = c.execute("SELECT id,staff_id,name FROM employees").fetchall()
        records = {int(r["employee_id"]): r for r in c.execute(
            "SELECT id,employee_id,net_salary,payment_status FROM payroll_records WHERE salary_month=?", (month,)).fetchall()}
    exact = {_norm(e["staff_id"]): e for e in employees}
    by_digits: dict[str, list] = {}
    for e in employees:
        by_digits.setdefault(_digits(e["staff_id"]), []).append(e)
    seen, out = set(), []
    for index, row in enumerate(rows):
        amount = row.get("total") if row.get("total") is not None else row.get("net")
        employee = exact.get(_norm(row["staff_id"]))
        if not employee and _digits(row["staff_id"]):
            candidates = by_digits.get(_digits(row["staff_id"]), [])
            employee = candidates[0] if len(candidates) == 1 else None
        item = {"index": index, "row": row, "amount": amount, "employee_id": None, "employee_name": "",
                "system_amount": None, "system_status": "", "action": "skip", "reason": "", "warning": ""}
        parts = [row.get(k) for k in ("gross", "absent_deduction", "overtime", "extra", "net")]
        if None not in parts and abs(parts[0] - parts[1] + parts[2] + parts[3] - parts[4]) > 1:
            item["warning"] = "Gross − Absent + OT + Extra ≠ Net; number মিলিয়ে দেখুন"
        if amount is not None and row.get("net") is not None and not 0 <= row["net"] - amount < 100:
            item["warning"] = (item["warning"] + "; " if item["warning"] else "") + "Total আর Net-এর পার্থক্য বেশি"
        if not employee:
            item["reason"] = "এই Staff ID-র employee নেই"
        elif int(employee["id"]) in seen:
            item["reason"] = "একই employee sheet-এ দুইবার"
        elif amount is None:
            item["reason"] = "Total Salary পড়া যায়নি"
        elif amount <= 0:
            item["reason"] = "টাকা ০, বাদ"
        else:
            record = records.get(int(employee["id"]))
            item.update(employee_id=int(employee["id"]), employee_name=employee["name"])
            if record:
                item.update(system_amount=float(record["net_salary"] or 0), system_status=record["payment_status"])
            if record and record["payment_status"] == "paid":
                item["reason"] = "আগেই Paid"
            else:
                item["action"] = "pay"
        if employee:
            seen.add(int(employee["id"]))
            item["employee_name"] = employee["name"]
        out.append(item)
    return out


# ------------------------------------------------------------------- apply

def _snapshot(row: dict, amount: float) -> dict:
    v = {k: float(row.get(k) or 0) for k in FIELDS}
    rounding = max(v["net"] - amount, 0) if v["net"] else 0
    gross = (v["gross"] or v["basic"] + v["night"] + v["friday"]) + v["overtime"] + v["extra"]
    total_deduction = v["absent_deduction"] + rounding
    return {"fixed_salary": v["basic"], "earned_basic_salary": v["basic"], "night_allowance": v["night"],
            "friday_allowance": v["friday"], "worked_duty_days": v["present"], "worked_duty_units": v["present"],
            "scheduled_duty_days": v["present"] + v["absent"] + v["leave"], "absent_days": v["absent"],
            "absent_duty_units": v["absent"], "paid_leave_days": v["leave"], "absent_deduction": v["absent_deduction"],
            "overtime_amount": v["overtime"], "overtime_hours": 0, "overtime_rate": 0, "bonus": v["extra"],
            "deduction": rounding, "late_minutes": 0, "late_deduction": 0, "advance_amount": 0, "fine_amount": 0,
            "gross_salary": gross, "total_deduction": total_deduction, "net_salary": amount,
            "salary_sheet_row": row}


def apply(import_id: int, picked: set[int], actor: str, method: str, reference: str) -> dict:
    from app.main import _log_payroll_change
    with get_db() as c:
        sheet = c.execute("SELECT * FROM salary_sheet_imports WHERE id=?", (import_id,)).fetchone()
        if not sheet:
            raise HTTPException(404, "Import not found")
        if sheet["applied_at"]:
            raise HTTPException(409, "This sheet was already applied")
        month = sheet["salary_month"]
        results, paid, total = [], 0, 0.0
        for item in plan(json.loads(sheet["rows_json"]), month):
            if item["action"] != "pay" or item["index"] not in picked:
                continue
            snap = _snapshot(item["row"], float(item["amount"]))
            before = c.execute("SELECT * FROM payroll_records WHERE employee_id=? AND salary_month=?",
                               (item["employee_id"], month)).fetchone()
            if before and before["payment_status"] == "paid":
                continue
            reason = (f"Salary sheet #{import_id}: system ৳{float(before['net_salary'] or 0):,.0f} → sheet ৳{item['amount']:,.0f}"
                      if before else f"Salary sheet #{import_id}: ৳{item['amount']:,.0f} (no system record)")
            values = (snap["fixed_salary"], snap["net_salary"], snap["overtime_amount"], snap["bonus"], snap["deduction"],
                      snap["gross_salary"], snap["total_deduction"], snap["absent_deduction"], snap["earned_basic_salary"],
                      snap["night_allowance"], snap["friday_allowance"], int(snap["worked_duty_days"]),
                      int(snap["absent_days"]), int(snap["paid_leave_days"]), json.dumps(snap, default=str),
                      method, reference, actor, reason[:500])
            if before:
                c.execute("UPDATE payroll_records SET fixed_salary=?,net_salary=?,overtime_amount=?,bonus=?,deduction=?,"
                          "gross_salary=?,total_deduction=?,absent_deduction=?,earned_basic_salary=?,night_allowance=?,"
                          "friday_allowance=?,worked_duty_days=?,absent_days=?,paid_leave_days=?,calculation_snapshot=?,"
                          "payment_method=?,payment_reference=?,locked_by=?,adjustment_reason=?,payment_status='paid',"
                          "finalized_at=COALESCE(finalized_at,CURRENT_TIMESTAMP),locked_at=COALESCE(locked_at,CURRENT_TIMESTAMP),"
                          "paid_at=CURRENT_TIMESTAMP,updated_by=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                          values + (actor, before["id"]))
                payroll_id = int(before["id"])
            else:
                c.execute("INSERT INTO payroll_records(fixed_salary,net_salary,overtime_amount,bonus,deduction,gross_salary,"
                          "total_deduction,absent_deduction,earned_basic_salary,night_allowance,friday_allowance,"
                          "worked_duty_days,absent_days,paid_leave_days,calculation_snapshot,payment_method,payment_reference,"
                          "locked_by,adjustment_reason,created_by,updated_by,employee_id,salary_month,overtime_mode,"
                          "payment_status,finalized_at,locked_at,paid_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,"
                          "'manual','paid',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)",
                          values + (actor, actor, item["employee_id"], month))
                payroll_id = int(c.execute("SELECT id FROM payroll_records WHERE employee_id=? AND salary_month=?",
                                           (item["employee_id"], month)).fetchone()["id"])
            _log_payroll_change(c, payroll_id, "paid", actor, reason)
            results.append({"payroll_id": payroll_id, "created": not before,
                            "before": {k: before[k] for k in before.keys()} if before else None})
            paid += 1
            total += float(item["amount"])
        c.execute("UPDATE salary_sheet_imports SET applied_at=?,applied_by=?,results_json=? WHERE id=?",
                  (int(time.time()), actor, json.dumps(results, default=str), import_id))
    return {"paid": paid, "total": total}


def undo(import_id: int, actor: str) -> int:
    from app.main import _log_payroll_change
    with get_db() as c:
        sheet = c.execute("SELECT * FROM salary_sheet_imports WHERE id=?", (import_id,)).fetchone()
        if not sheet or not sheet["applied_at"] or sheet["undone_at"]:
            raise HTTPException(409, "Nothing to undo")
        if time.time() - int(sheet["applied_at"]) > UNDO_HOURS * 3600:
            raise HTTPException(409, f"Undo is only possible within {UNDO_HOURS} hours")
        restored = 0
        for result in json.loads(sheet["results_json"] or "[]"):
            payroll_id = int(result["payroll_id"])
            if result["created"]:
                c.execute("DELETE FROM payroll_change_logs WHERE payroll_id=?", (payroll_id,))
                c.execute("DELETE FROM payroll_records WHERE id=?", (payroll_id,))
            else:
                before = result["before"]
                columns = [k for k in before if k not in ("id", "employee_id", "salary_month")]
                c.execute(f"UPDATE payroll_records SET {','.join(f'{k}=?' for k in columns)} WHERE id=?",
                          tuple(before[k] for k in columns) + (payroll_id,))
                _log_payroll_change(c, payroll_id, "sheet_undo", actor, f"Salary sheet #{import_id} undone")
            restored += 1
        c.execute("UPDATE salary_sheet_imports SET undone_at=?,undone_by=? WHERE id=?", (int(time.time()), actor, import_id))
    return restored


# -------------------------------------------------------------------- pages

def _money(value) -> str:
    return "—" if value is None else f"৳{float(value):,.0f}"


def _month_options(selected: str) -> str:
    from app.services import now_local
    now = now_local().date().replace(day=1)
    options, year, month = [], now.year, now.month
    for _ in range(12):
        key = f"{year:04d}-{month:02d}"
        label = f"{MONTHS[month - 1].title()} {year}"
        options.append(f"<option value='{key}' {'selected' if key == selected else ''}>{label}</option>")
        month -= 1
        if month == 0:
            year, month = year - 1, 12
    return "".join(options)


def _default_month() -> str:
    from app.services import now_local
    today = now_local().date()
    last = today.replace(day=1)
    year, month = (last.year, last.month - 1) if last.month > 1 else (last.year - 1, 12)
    return f"{year:04d}-{month:02d}"


@router.get("/payroll/sheet", response_class=HTMLResponse)
def sheet_page(request: Request, error: str = "", done: str = ""):
    from app.main import require_permission, layout
    require_permission(request, "payroll_manage")
    with get_db() as c:
        history = c.execute("SELECT id,salary_month,file_name,source,created_by,created_at,applied_at,applied_by,undone_at,results_json "
                            "FROM salary_sheet_imports ORDER BY id DESC LIMIT 15").fetchall()
    note = ""
    if error:
        note = f"<div class='card' style='border-color:#e8b4ad;background:#fdf0ee'><b>{escape(error)}</b></div><div class='section-gap'></div>"
    elif done:
        note = f"<div class='card' style='border-color:#a6e8c8;background:#eefaf3'><b>{escape(done)}</b></div><div class='section-gap'></div>"
    rows = []
    for h in history:
        count = len(json.loads(h["results_json"] or "[]"))
        if h["undone_at"]:
            state = "<span class='status bad'>Undone</span>"
        elif h["applied_at"]:
            state = f"<span class='status ok'>{count} paid</span>"
        else:
            state = f"<a class='btn secondary' href='/payroll/sheet/{h['id']}'>Review</a>"
        undo_btn = ""
        if h["applied_at"] and not h["undone_at"] and time.time() - int(h["applied_at"]) < UNDO_HOURS * 3600 \
                and request.session.get("role") in ("admin", "super_admin"):
            undo_btn = (f"<form method='post' action='/payroll/sheet/{h['id']}/undo' style='display:inline'>"
                        "<button class='btn secondary' onclick=\"return confirm('Undo this import? Every record goes back to before.')\">Undo</button></form>")
        rows.append(f"<tr><td>#{h['id']}</td><td><b>{escape(h['salary_month'])}</b></td>"
                    f"<td><a href='/payroll/sheet/{h['id']}/file'>{escape(h['file_name'])}</a><div class='sub'>{'Excel' if h['source'] == 'excel' else 'AI read'}</div></td>"
                    f"<td>{escape(h['applied_by'] or h['created_by'] or '')}</td><td>{state} {undo_btn}</td></tr>")
    body = f"""{note}<div class='hero'><div><h2>Salary sheet upload</h2>
      <div class='sub'>At month end, upload the salary sheet you already make. Check the preview, then confirm, and everyone on it is marked Paid with the sheet's Total Salary.</div>
      <div class='sub' style='margin-top:6px'>Excel is read exactly. A PDF or photo is read by AI (Gemini), so check every row in the preview.</div></div></div>
    <div class='card'><form method='post' action='/payroll/sheet' enctype='multipart/form-data'>
      <div class='grid'><div><label>Salary month</label><select name='month'>{_month_options(_default_month())}</select></div>
      <div><label>Sheet file (.xlsx, PDF, JPG or PNG)</label><input type='file' name='sheet' required accept='.xlsx,.pdf,image/*'></div></div>
      <button class='btn'>Read the sheet</button></form></div>
    <div class='section-gap'></div><div class='card' style='overflow:auto'><h3>Uploaded sheets</h3>
    <table><thead><tr><th></th><th>Month</th><th>File</th><th>By</th><th>Status</th></tr></thead>
    <tbody>{''.join(rows) or "<tr><td colspan=5 class='sub'>No sheet uploaded yet.</td></tr>"}</tbody></table></div>"""
    return layout("Salary sheet upload", body, request, "salarysheet")


@router.post("/payroll/sheet")
async def sheet_upload(request: Request, month: str = Form(...), sheet: UploadFile = File(...)):
    import asyncio
    from app.main import require_permission, _payroll_actor
    from app.ai_insights import AIUnavailable
    require_permission(request, "payroll_manage")
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise HTTPException(400, "Invalid month")
    content = await sheet.read(MAX_FILE_BYTES + 1)
    name = (sheet.filename or "sheet").rsplit("/", 1)[-1][:120]
    if not content or len(content) > MAX_FILE_BYTES:
        return RedirectResponse("/payroll/sheet?error=" + _q("File খালি বা ১২ MB-এর বেশি।"), 303)
    lower = name.lower()
    if lower.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        mime, source = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "excel"
    elif lower.endswith(".pdf") or content[:4] == b"%PDF":
        mime, source = "application/pdf", "ai"
    elif content[:3] == b"\xff\xd8\xff":
        mime, source = "image/jpeg", "ai"
    elif content[:8] == b"\x89PNG\r\n\x1a\n":
        mime, source = "image/png", "ai"
    elif content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        mime, source = "image/webp", "ai"
    else:
        return RedirectResponse("/payroll/sheet?error=" + _q("শুধু .xlsx, PDF, JPG বা PNG দিন।"), 303)
    try:
        parsed = read_excel(content) if source == "excel" else await asyncio.to_thread(read_with_ai, content, mime)
    except (ValueError, AIUnavailable) as exc:
        return RedirectResponse("/payroll/sheet?error=" + _q(str(exc)), 303)
    if not parsed["rows"]:
        return RedirectResponse("/payroll/sheet?error=" + _q("Sheet-এ কোনো employee row পাওয়া যায়নি।"), 303)
    with get_db() as c:
        c.execute("INSERT INTO salary_sheet_imports(salary_month,file_name,file_type,file_data,source,rows_json,printed_total,"
                  "month_label,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                  (month, name, mime, base64.b64encode(content).decode(), source, json.dumps(parsed["rows"]),
                   parsed["printed_total"], parsed["month_label"], _payroll_actor(request), int(time.time())))
        import_id = c.execute("SELECT MAX(id) AS id FROM salary_sheet_imports").fetchone()["id"]
    logger.info("Salary sheet uploaded id=%s month=%s rows=%s source=%s", import_id, month, len(parsed["rows"]), source)
    return RedirectResponse(f"/payroll/sheet/{import_id}", 303)


def _q(message: str) -> str:
    from urllib.parse import quote
    return quote(message[:300])


@router.get("/payroll/sheet/{import_id}", response_class=HTMLResponse)
def sheet_review(request: Request, import_id: int):
    from app.main import require_permission, layout
    require_permission(request, "payroll_manage")
    with get_db() as c:
        sheet = c.execute("SELECT * FROM salary_sheet_imports WHERE id=?", (import_id,)).fetchone()
    if not sheet:
        raise HTTPException(404, "Import not found")
    if sheet["applied_at"]:
        return RedirectResponse("/payroll/sheet", 303)
    month = sheet["salary_month"]
    items = plan(json.loads(sheet["rows_json"]), month)
    payable = [i for i in items if i["action"] == "pay"]
    sheet_sum = sum(float(i["amount"] or 0) for i in items if i["amount"])
    pay_sum = sum(float(i["amount"]) for i in payable)
    warnings = []
    printed = sheet["printed_total"]
    if printed is not None and abs(float(printed) - sheet_sum) > 1:
        warnings.append(f"Sheet-এর নিচের মোট {_money(printed)}, কিন্তু row-গুলোর যোগফল {_money(sheet_sum)}। কোনো row ভুল পড়া হয়েছে কিনা দেখুন।")
    month_number = int(month[5:])
    label = (sheet["month_label"] or "").strip().lower()
    if label and not MONTHS[month_number - 1].startswith(label[:3]):
        warnings.append(f"Sheet-এ মাস লেখা '{sheet['month_label']}', কিন্তু আপনি বেছেছেন {MONTHS[month_number - 1].title()}। মাস ঠিক আছে তো?")
    flagged = sum(1 for i in items if i["warning"])
    if flagged:
        warnings.append(f"{flagged}টি row-এর হিসাব নিজের মধ্যে মেলেনি। হলুদ row-গুলো sheet-এর সাথে মিলিয়ে দেখুন।")
    if sheet["source"] == "ai":
        warnings.append("এই sheet AI পড়েছে। Confirm করার আগে টাকার অঙ্কগুলো মূল sheet-এর সাথে মিলিয়ে নিন।")
    warn_html = "".join(f"<div class='card' style='border-color:#eed29a;background:#fdf7e8;margin-bottom:10px'>{escape(w)}</div>" for w in warnings)

    def row_html(i):
        r = i["row"]
        pay = i["action"] == "pay"
        diff = ""
        if pay and i["system_amount"] is not None and abs(i["system_amount"] - float(i["amount"])) >= 1:
            diff = f"<div class='sub'>পার্থক্য {_money(float(i['amount']) - i['system_amount'])}</div>"
        system = (f"{_money(i['system_amount'])}<div class='sub'>{escape(i['system_status'])}</div>"
                  if i["system_amount"] is not None else "<span class='sub'>No record</span>")
        tick = (f"<input type='checkbox' name='pick' value='{i['index']}' checked aria-label='Pay {escape(r['name'])}'>" if pay else "")
        state = "<span class='status ok'>Will be paid</span>" if pay else f"<span class='status warn'>{escape(i['reason'])}</span>"
        style = " style='background:#fdf7e8'" if i["warning"] else ""
        return (f"<tr{style}><td>{tick}</td><td><b>{escape(r['name'] or '—')}</b><div class='sub'>{escape(r['staff_id'])}"
                f"{' → ' + escape(i['employee_name']) if i['employee_name'] and i['employee_name'] != r['name'] else ''}</div>"
                f"{'<div class=sub>' + escape(i['warning']) + '</div>' if i['warning'] else ''}</td>"
                f"<td>{'' if r.get('present') is None else int(r['present'])}</td><td>{_money(r.get('net'))}</td>"
                f"<td><b>{_money(i['amount'])}</b></td><td>{system}{diff}</td><td>{state}</td></tr>")

    body = f"""<div class='hero'><div><h2>Check the sheet: {escape(MONTHS[month_number - 1].title())} {month[:4]}</h2>
      <div class='sub'>{escape(sheet['file_name'])}. {len(items)} rows read, {len(payable)} will be marked Paid, total {_money(pay_sum)}.</div>
      <div class='sub'>The sheet's Total Salary becomes the final paid amount. The system's figure is kept in each record's history.</div></div>
      <a class='btn secondary' href='/payroll/sheet/{import_id}/file' target='_blank'>Open the uploaded file</a></div>
    {warn_html}
    <form method='post' action='/payroll/sheet/{import_id}/apply'>
    <div class='card' style='overflow:auto'><table><thead><tr><th></th><th>Employee</th><th>Present</th><th>Net</th><th>Total (paid)</th>
      <th>System</th><th></th></tr></thead><tbody>{''.join(row_html(i) for i in items)}</tbody></table></div>
    <div class='section-gap'></div><div class='card'><div class='grid'>
      <div><label>Payment method</label><select name='method'><option>Cash</option><option>bKash</option><option>Nagad</option><option>Bank</option></select></div>
      <div><label>Reference</label><input name='reference' required maxlength='120' value='Salary sheet {escape(month)}'></div></div>
      <div class='actions'><button class='btn' onclick="return confirm('Mark the ticked employees as Paid?')">Confirm: mark ticked as Paid</button>
      <a class='btn secondary' href='/payroll/sheet'>Cancel</a></div>
      <div class='sub'>An Admin can undo this for {UNDO_HOURS} hours.</div></div></form>"""
    return layout("Check salary sheet", body, request, "salarysheet")


@router.post("/payroll/sheet/{import_id}/apply")
async def sheet_apply(request: Request, import_id: int):
    from app.main import require_permission, _payroll_actor, audit
    require_permission(request, "payroll_manage")
    form = await request.form()
    picked = {int(v) for v in form.getlist("pick") if str(v).isdigit()}
    method = str(form.get("method") or "Cash").strip()[:40] or "Cash"
    reference = str(form.get("reference") or "").strip()[:120]
    if not reference:
        raise HTTPException(400, "Reference is required")
    result = apply(import_id, picked, _payroll_actor(request), method, reference)
    audit(request, "salary_sheet_apply", "payroll", str(import_id), f"{result['paid']} paid, ৳{result['total']:,.0f}")
    return RedirectResponse("/payroll/sheet?done=" + _q(f"{result['paid']} জনের বেতন Paid হয়েছে, মোট ৳{result['total']:,.0f}।"), 303)


@router.post("/payroll/sheet/{import_id}/undo")
def sheet_undo(request: Request, import_id: int):
    from app.main import require_permission, _payroll_actor, audit
    require_permission(request, "payroll_manage")
    if request.session.get("role") not in ("admin", "super_admin"):
        raise HTTPException(403, "Admin access required")
    restored = undo(import_id, _payroll_actor(request))
    audit(request, "salary_sheet_undo", "payroll", str(import_id), f"{restored} records restored")
    return RedirectResponse("/payroll/sheet?done=" + _q(f"Import #{import_id} undo হয়েছে: {restored}টি record আগের অবস্থায়।"), 303)


@router.get("/payroll/sheet/{import_id}/file")
def sheet_file(request: Request, import_id: int):
    from app.main import require_permission
    require_permission(request, "payroll_view")
    with get_db() as c:
        sheet = c.execute("SELECT file_name,file_type,file_data FROM salary_sheet_imports WHERE id=?", (import_id,)).fetchone()
    if not sheet:
        raise HTTPException(404, "Import not found")
    disposition = "inline" if sheet["file_type"].startswith(("image/", "application/pdf")) else "attachment"
    safe = re.sub(r"[^\w.\-]", "_", sheet["file_name"])
    return Response(base64.b64decode(sheet["file_data"]), media_type=sheet["file_type"],
                    headers={"Content-Disposition": f"{disposition}; filename={safe}", "Cache-Control": "no-store"})
